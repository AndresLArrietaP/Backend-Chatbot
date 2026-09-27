"""
Genera el documento Word 'Seguimiento de Pedidos — Partes Interesadas'.

Reúne, en dos tablas, los pedidos de las reuniones de las últimas 2 semanas y cómo/dónde se
cumplieron:
  - Tabla 1: INVERTEX 4 — reuniones técnicas 17/07 y 21/07/2026 (juntas). Se dejan fuera los
    pedidos aún pendientes/a futuro (predefinidos por evento, payload desde sprung, leyenda de
    estados, decodificar digitales, módulo de interfaces).
  - Tabla 2: KomfIA — pedidos de gerencia (los del cuadro de seguimiento).

Estilo reutilizado de tools/agregar_comparativa.py (mismos helpers de docx). Salida a las dos
carpetas 'gerencia' (INVERTEX v3/refs/gerencia y Backend-Chatbot/docs/gerencia), como el acta.
"""
import shutil

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x1F, 0x49, 0x7D)
# Sombreado de la columna Estado (como el cuadro de la imagen).
EST_FILL = {'ok': 'E4F2E4', 'wip': 'FBF0D8', 'pend': 'ECECEC'}
EST_FONT = {'ok': RGBColor(0x1E, 0x6A, 0x30), 'wip': RGBColor(0x8A, 0x5A, 0x14),
            'pend': RGBColor(0x66, 0x66, 0x66)}

SALIDAS = [
    r'c:\Users\Usuario\Pictures\PYTHON\INVERTEX v3\refs\gerencia\SEGUIMIENTO_Pedidos_partes_interesadas.docx',
    r'c:\Users\Usuario\Pictures\PYTHON\Backend-Chatbot\docs\gerencia\SEGUIMIENTO_Pedidos_partes_interesadas.docx',
]


# ── helpers de estilo (calco de tools/agregar_comparativa.py) ───────────────────
def _shade(cell, hex_fill):
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hex_fill)
    cell._tc.get_or_add_tcPr().append(shd)


def _set_cell(cell, texto, size=9, bold=False, color=None, fill=None, align_center=False):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(1); p.paragraph_format.space_after = Pt(1)
    if align_center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(str(texto)); r.font.size = Pt(size); r.bold = bold
    if color:
        r.font.color.rgb = color
    if fill:
        _shade(cell, fill)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_table(doc, headers, rows, col_widths, est_col=None):
    """Tabla con cabecera navy (texto blanco) y, si `est_col`, sombreado de estado en esa columna.
    Cada fila es (val0, val1, …); si hay est_col, el valor de esa columna es (texto, tipo)."""
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = 'Table Grid'
    for i, h in enumerate(headers):                    # cabecera
        _set_cell(t.rows[0].cells[i], h, size=9.5, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF),
                  fill='1F497D', align_center=True)
    for ri, row in enumerate(rows):                    # datos
        for ci, val in enumerate(row):
            cell = t.rows[ri + 1].cells[ci]
            if ci == est_col:
                texto, tipo = val
                _set_cell(cell, texto, size=9, bold=True, color=EST_FONT[tipo], fill=EST_FILL[tipo])
            else:
                _set_cell(cell, val, size=9, align_center=(ci == 0))
    for i, w in enumerate(col_widths):
        for row in t.rows:
            row.cells[i].width = Inches(w)
    doc.add_paragraph()
    return t


# ── documento ───────────────────────────────────────────────────────────────────
doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = sec.page_height, sec.page_width      # letter horizontal
for m in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
    setattr(sec, m, Inches(0.6))
doc.styles['Normal'].font.name = 'Calibri'
doc.styles['Normal'].font.size = Pt(10)

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Seguimiento de Pedidos — Partes Interesadas')
r.font.size = Pt(18); r.bold = True; r.font.color.rgb = NAVY
p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r2 = p2.add_run('INVERTEX 4 (reuniones técnicas 17 y 21/07/2026)  ·  KomfIA (gerencia)  ·  '
                'Generado 23/07/2026')
r2.font.size = Pt(10); r2.italic = True; r2.font.color.rgb = RGBColor(0x5A, 0x5A, 0x5A)
doc.add_paragraph()

# ── Tabla 1: INVERTEX 4 (17/07 + 21/07) ─────────────────────────────────────────
h = doc.add_heading('INVERTEX 4 — Reuniones 17/07 y 21/07/2026', level=1)
h.runs[0].font.size = Pt(13); h.runs[0].font.color.rgb = NAVY

inv_headers = ['#', 'Reunión', 'Pedido', 'Cómo / dónde lo cumplimos', 'Estado']
inv_rows = [
    ('1', '17/07', 'Detección de patrón operativo: juntar los datapacks del evento y hallar la '
     'similitud de condición operativa (velocidad, pendiente, carga/sprung); solo con data Invertex.',
     'Módulo de patrón (core/patron.py) con aviso de representatividad de la muestra.',
     ('✔  Cumplido', 'ok')),
    ('2', '17/07', 'Filtro por evento en Análisis por Evento (antes salía solo el general).',
     'Filtro de camión/evento en la vista de Evento (ui/paginas/evento.py).',
     ('✔  Cumplido', 'ok')),
    ('3', '21/07', 'Volver los textos del datapack al inglés original (no traducirlos al español).',
     'Re-etiquetado al estándar inglés del manual GEK-91842A en datapack_parametros.json y datalr.py.',
     ('✔  Cumplido', 'ok')),
    ('4', '21/07', 'Ruta y Perfil / LR: mostrar más parámetros (no solo 5).',
     'Rediseño a formato Datapack: búsqueda + multi-eje con todas las columnas del LR (ruta_lr.py).',
     ('✔  Cumplido', 'ok')),
    ('5', '21/07', 'Poder ocultar / activar variables en el gráfico.',
     'Selección de parámetros con casillas en Ruta/Perfil y en el panel del Datapack.',
     ('✔  Cumplido', 'ok')),
    ('6', '21/07', 'Mapa de Calor: quitar lo que resta espacio sin aportar valor (el ranking de barras).',
     'Se quitó el ranking; queda solo el mapa de calor Mes × Camión (mapa_calor.py).',
     ('✔  Cumplido', 'ok')),
    ('7', '21/07', 'Clasificar los eventos: no-power (críticos) vs sys-event (informativos).',
     'Analizado en el datapack (tipos Log / Log+ / ----): NO trae un flag de encendido; la señal real '
     'está en la telemetría. No se fabricó etiqueta (queda por definir la regla con Carlos).',
     ('◐  Analizado', 'wip')),
    ('8', '21/07', 'Generador de Excel / reporte de evento (elegir qué incluir).',
     'Barrido de evento a Excel: Resumen, Por equipo, Por mes, Mapa de calor y Georreferencia (barrido.py).',
     ('✔  Cumplido', 'ok')),
    ('9', '21/07', 'Diseño más minimalista (menos saturación de color).',
     'Dirección de rediseño del vestíbulo/tarjetas; en curso.',
     ('◐  En curso', 'wip')),
]
add_table(doc, inv_headers, inv_rows, col_widths=[0.4, 0.8, 3.2, 3.7, 1.35], est_col=4)

# ── Tabla 2: KomfIA (gerencia) ──────────────────────────────────────────────────
h2 = doc.add_heading('KomfIA — Pedidos de gerencia', level=1)
h2.runs[0].font.size = Pt(13); h2.runs[0].font.color.rgb = NAVY

kom_headers = ['#', 'Pedido de gerencia', 'Cómo / dónde lo cumplimos', 'Estado']
kom_rows = [
    ('1', 'Barrido en matriz compacta por componente.',
     'Formatos (caso A) + KomfIA_central (barrido PASO 2).',
     ('✔  Cumplido y desplegado', 'ok')),
    ('2', 'Acumulado / Σvida: sumatoria total del metal en la vida del componente.',
     'vw_TendenciaElemento (cols Acumulado / NmAcum) + columna Σvida en el resumen; '
     'Formatos / Esquema · BLOQUE 28.',
     ('✔  Cumplido y validado', 'ok')),
    ('3', 'Límites visibles en el barrido: para qué componente y con qué LP/LC (cuadro componente × metal).',
     'Detalle con (LP/LC) + LimObs / Limites en las vistas + cuadro en Formatos · BLOQUES 30/31.',
     ('✔  Cumplido y desplegado', 'ok')),
    ('4', 'Velocidad (Antamina ~2 min).',
     'Índices del DBA (14/07) → queries ~0.4-0.7 s.',
     ('◐  Índices puestos; render = optim. pendiente', 'wip')),
    ('5', 'Exportar PDF / Excel / Word (no urgente).',
     'Requiere flujo Power Automate.',
     ('⏳  Pendiente (futuro)', 'pend')),
    ('6', 'KPIs por correo + reuniones cortas semanales (no urgente).',
     'Flujo aparte / metodología de seguimiento.',
     ('⏳  Pendiente', 'pend')),
]
add_table(doc, kom_headers, kom_rows, col_widths=[0.4, 3.5, 4.25, 1.6], est_col=3)

# ── guardar + sincronizar ────────────────────────────────────────────────────────
doc.save(SALIDAS[0])
for destino in SALIDAS[1:]:
    shutil.copyfile(SALIDAS[0], destino)
print('Documento generado y sincronizado en:')
for s in SALIDAS:
    print('  -', s)
